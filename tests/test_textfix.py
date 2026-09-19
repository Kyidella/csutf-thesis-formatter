"""文本层修复的测试。

这一层是唯一"改字"的，所以测试的重心和前面几层不同：
  1. 该改的改了，且改的位置对（下标映射不能偏）
  2. **w:instrText 一个字节都不能动**——Zotero 的 JSON 在里面，改坏不报错
  3. 域内的字不改，域外的照改（这条是政策，见 textfix 开头）
  4. run / w:t 节点数量不变（只改文字，不新增不合并）
  5. 首尾空格要补 xml:space，否则 Word 把空格吃掉
  6. 检查器报的 = 格式化器改的

注：异常字符一律写成 \\uXXXX 转义。NBSP、零宽空格、私有区字符都是**看不见**的，
直接写成字面量的话，谁review都看不出这一行到底在测什么。
"""

import pytest
import yaml
from docx import Document
from docx.oxml.ns import qn

from classifier import classify
from conftest import RULES
from docmodel import W, load
from textfix import (TextChange, caption_gap_pos, fix_text, flatten, protected,
                     punct_fixes)

pytestmark = pytest.mark.skipif(not RULES.exists(), reason="规则文件缺失")

XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
NBSP = chr(0x00A0)        # 不换行空格（WPS 里显示成 °）
PUA_SPACE = chr(0xE5E5)   # 私有区字符，样本里顶替空格用
PUA_JUNK = chr(0xE000)    # 其他私有区字符，属于垃圾
ZWSP = chr(0x200B)        # 零宽空格
REPL = chr(0xFFFD)        # 替换字符
MICRO_BAD = chr(0x00B5)   # 微符号 µ
MICRO_GOOD = chr(0x03BC)  # 希腊字母 μ


@pytest.fixture(scope="module")
def rules():
    return yaml.safe_load(open(RULES, encoding="utf-8"))


# ---- 造文档的小工具 ----
def make_doc(tmp_path, *paragraphs, name="t.docx"):
    """每项是一个字符串（单 run）或字符串列表（多 run）。"""
    doc = Document()
    for spec in paragraphs:
        p = doc.add_paragraph()
        for s in ([spec] if isinstance(spec, str) else spec):
            p.add_run(s)
    path = tmp_path / name
    doc.save(path)
    return path


def fldChar(run, kind):
    return run._element.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): kind})


def add_field(p, instr=" ADDIN ZOTERO_ITEM CSL_CITATION {} ", display="[1]"):
    """往段落里塞一个完整的引用域（begin / instrText / separate / 结果 / end）。"""
    for kind in ("begin", None, "separate"):
        r = p.add_run()
        if kind is None:
            node = r._element.makeelement(qn("w:instrText"), {})
            node.text = instr
            r._element.append(node)
        else:
            r._element.append(fldChar(r, kind))
    p.add_run(display)
    r = p.add_run()
    r._element.append(fldChar(r, "end"))


def run_fix(path, rules):
    """跑一遍文本层修复，返回 (doc, changes, notes)。"""
    doc = load(path)
    changes, notes = [], []
    fix_text(doc, classify(doc), rules, changes, notes)
    return doc, changes, notes


def after(doc):
    """修复后的段落文字（口径与 docmodel._text_of 一致）。"""
    return [flatten(p.element)[0] for p in doc.paragraphs]


# --------------------------------------------------------------------------
# 位置映射与域保护：整层都建立在这两个函数上，先验它们
# --------------------------------------------------------------------------
def test_位置映射与段落文本一致(tmp_path):
    """flatten 的下标必须和 docmodel 算出来的 p.text 一字符不差。

    对不上就会改到隔壁的字——这是这一层最危险的 bug。
    """
    path = make_doc(tmp_path, ["表2.1", "洗脱梯度"], "含制表符\t的段落")
    doc = load(path)
    for p in doc.paragraphs:
        assert flatten(p.element)[0] == p.text


def test_域内的节点被标为受保护(tmp_path):
    path = make_doc(tmp_path, "前")
    d = Document(path)
    p = d.paragraphs[0]
    add_field(p)
    p.add_run("后")
    d.save(path)

    doc = load(path)
    el = doc.paragraphs[0].element
    prot = protected(el)
    ts = list(el.iter(f"{W}t"))
    # 顺序：前 / [1] / 后 —— 只有中间的 [1] 在域里
    assert [t in prot for t in ts] == [False, True, False]


def test_fldSimple里的节点也算域内(tmp_path):
    """有些写入器用 w:fldSimple 整元素成域，不靠 fldChar 配对。"""
    path = make_doc(tmp_path, "正文")
    d = Document(path)
    p = d.paragraphs[0]
    simple = p._element.makeelement(qn("w:fldSimple"), {qn("w:instr"): " TIME "})
    r = simple.makeelement(qn("w:r"), {})
    t = simple.makeelement(qn("w:t"), {})
    t.text = "12:00"
    r.append(t)
    simple.append(r)
    p._element.insert(0, simple)
    d.save(path)

    doc = load(path)
    assert protected(doc.paragraphs[0].element)


# --------------------------------------------------------------------------
# 一、题注缺空格
# --------------------------------------------------------------------------
def test_题注缺空格_同节点内插入(tmp_path, rules):
    path = make_doc(tmp_path, "图2.1油茶果实发育")
    doc, changes, _ = run_fix(path, rules)
    assert after(doc)[0] == "图2.1 油茶果实发育"
    assert any("题注" in c.what for c in changes)
    # 同一个 w:t 里从中间插空格，不在首尾，不需要 xml:space
    assert flatten(doc.paragraphs[0].element)[1][0][0].get(XML_SPACE) is None


def test_题注缺空格_跨节点挂到图号末尾并加preserve(tmp_path, rules):
    """图号和图名分在两个 w:t 里时，空格挂到图号那个节点尾部。

    这时节点末位成了空格，必须补 xml:space="preserve"，
    否则 Word 渲染时把空格吃掉，题注看着还是没改。
    实测某样本论文里有 2 处正是这种跨节点的。
    """
    path = make_doc(tmp_path, ["表3.1", "洗脱梯度"])
    doc, _, _ = run_fix(path, rules)
    ts = list(doc.paragraphs[0].element.iter(f"{W}t"))
    assert ts[0].text == "表3.1 "
    assert ts[0].get(XML_SPACE) == "preserve"
    assert ts[1].text == "洗脱梯度"
    assert after(doc)[0] == "表3.1 洗脱梯度"


def test_题注空格数由规则决定(tmp_path, rules):
    """把 gap_after_number 改成 2，就该插两个空格——不能再写死在代码里。

    这条是审查发现的：规则里写着 gap_after_number，代码却硬编码一个空格，
    改规则不会生效，看规则的人会以为它在起作用。
    """
    import copy

    patched = copy.deepcopy(rules)
    patched["styles"]["figure_caption"]["gap_after_number"] = 2
    path = make_doc(tmp_path, "图2.1油茶果实发育")
    doc, changes, _ = run_fix(path, patched)
    assert after(doc)[0] == "图2.1  油茶果实发育", "应插两个空格"
    assert any("插入 2 个半角空格" in c.new for c in changes)

    # 表题取自己那条，互不影响
    path2 = make_doc(tmp_path, ["表3.1", "洗脱梯度"], name="t2.docx")
    doc2, _, _ = run_fix(path2, patched)
    assert after(doc2)[0] == "表3.1 洗脱梯度", "表题仍是 1 个空格"


def test_题注已有空格不动(tmp_path, rules):
    path = make_doc(tmp_path, "图2.1 油茶果实发育")
    doc, changes, _ = run_fix(path, rules)
    assert after(doc)[0] == "图2.1 油茶果实发育"
    assert not changes


def test_题注没有图名时跳过并说明(tmp_path, rules):
    """整段只有"图2.1"：检查器会报缺空格，但没地方插，只能跳过并说明。"""
    path = make_doc(tmp_path, "图2.1")
    doc, _, notes = run_fix(path, rules)
    assert after(doc)[0] == "图2.1"
    assert any("没有图名" in n for n in notes)


def test_域内的题注不改(tmp_path, rules):
    """图号落在引用域里（罕见但会发生），只报不改。"""
    path = make_doc(tmp_path, "占位")
    d = Document(path)
    p = d.paragraphs[0]
    p.runs[0].text = ""
    add_field(p, display="图2.1")
    p.add_run("油茶果实")
    d.save(path)

    doc, changes, notes = run_fix(path, rules)
    assert not changes, "域内的题注不该被改"
    assert any("引用域" in n for n in notes)


# --------------------------------------------------------------------------
# 二、异常字符
# --------------------------------------------------------------------------
def test_NBSP换半角空格(tmp_path, rules):
    path = make_doc(tmp_path, "Camellia oleifera" + NBSP + "Abel. 是油茶")
    doc, changes, _ = run_fix(path, rules)
    assert after(doc)[0] == "Camellia oleifera Abel. 是油茶"
    assert any("NBSP" in c.what for c in changes)


def test_NBSP在词尾时补preserve(tmp_path, rules):
    """换出来的空格落在节点末尾，不加 preserve 会被 Word 吃掉，
    前后两个单词会被粘在一起——这是实测踩到的。"""
    path = make_doc(tmp_path, ["genes (e.g.," + NBSP, "CoEXPA2) expressed"])
    doc, _, _ = run_fix(path, rules)
    node = list(doc.paragraphs[0].element.iter(f"{W}t"))[0]
    assert node.text == "genes (e.g., "
    assert node.get(XML_SPACE) == "preserve"


def test_NBSP在词首时补preserve(tmp_path, rules):
    """换出来的空格落在节点开头，同样要保（首尾是两个独立分支）。"""
    path = make_doc(tmp_path, ["Camellia oleifera", NBSP + "Abel. 是油茶"])
    doc, _, _ = run_fix(path, rules)
    node = list(doc.paragraphs[0].element.iter(f"{W}t"))[1]
    assert node.text == " Abel. 是油茶"
    assert node.get(XML_SPACE) == "preserve"


def test_原有的裸首空格不被顺手补上():
    """_set_text 只补"我们新造出来的"首尾空格。

    原先就是裸首空格的（Word 一直当它不存在），补上 preserve 会改变别处的
    渲染结果——这种顺手改变别处的事，这一层一件都不做。

    直接测函数：python-docx 写字时会自动带上 preserve，造不出这种节点。
    """
    from lxml import etree

    from textfix import _set_text

    ns = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'

    a = etree.fromstring(f"<w:t {ns}></w:t>")
    _set_text(a, " 新造出来的首空格")
    assert a.get(XML_SPACE) == "preserve"

    b = etree.fromstring(f"<w:t {ns}> 本来就有</w:t>")
    _set_text(b, " 本来就有 改过")
    assert b.get(XML_SPACE) is None, "原先就有的裸空格不该被补 preserve"


def test_私有区字符被删除(tmp_path, rules):
    path = make_doc(tmp_path, "第一章" + PUA_JUNK + "绪论")
    doc, changes, _ = run_fix(path, rules)
    assert after(doc)[0] == "第一章绪论"
    assert any("私有区" in c.what for c in changes)


def test_当空格的私有区字符换成空格而不是删掉(tmp_path, rules):
    """U+E5E5 在样本里顶替空格用，删掉会让标题少一个间隔。

    证据：六个一级标题里，第二~六章都是"第二章 ＋ 空格 ＋ 标题"，
    只有第一章那个位置是 U+E5E5。规则里为它单列一条，排在任何通用区间之前。
    """
    path = make_doc(tmp_path, "第一章" + PUA_SPACE + "绪论")
    doc, changes, _ = run_fix(path, rules)
    assert after(doc)[0] == "第一章 绪论", "该换成空格，不是删掉"
    assert any("U+E5E5" in c.what for c in changes)


def test_零宽空格被删除(tmp_path, rules):
    path = make_doc(tmp_path, "正文" + ZWSP + "里有个零宽空格")
    doc, _, _ = run_fix(path, rules)
    assert after(doc)[0] == "正文里有个零宽空格"


def test_替换字符只报不改(tmp_path, rules):
    """U+FFFD 说明原稿这里本来就丢了字，自动删掉等于把线索也删了。"""
    path = make_doc(tmp_path, "数据" + REPL + "缺失")
    doc, changes, _ = run_fix(path, rules)
    assert after(doc)[0] == "数据" + REPL + "缺失"
    assert not changes, "替换字符不该被自动处理"


def test_微符号统一为首选(tmp_path, rules):
    path = make_doc(tmp_path, "进样量 10 " + MICRO_BAD + "L，另处作 " + MICRO_GOOD + "m")
    doc, changes, _ = run_fix(path, rules)
    assert after(doc)[0] == "进样量 10 " + MICRO_GOOD + "L，另处作 " + MICRO_GOOD + "m"
    assert any("微符号" in c.what for c in changes)


def test_instrText里的异常字符绝不修改(tmp_path, rules):
    """Zotero 把引用数据以 JSON 存在 w:instrText 里，那里面本来就有 NBSP。

    盲扫段落文本删字符，第一个受害的就是它，而且改坏了不报错。
    """
    path = make_doc(tmp_path, "正文")
    d = Document(path)
    add_field(d.paragraphs[0],
              instr=' ADDIN ZOTERO_ITEM CSL_CITATION {"id":' + NBSP + "1} ")
    d.paragraphs[0].add_run("尾巴" + NBSP + "也有")
    d.save(path)

    doc, _, _ = run_fix(path, rules)
    instr = [n.text for n in doc.root.iter(f"{W}instrText")]
    assert any(NBSP in s for s in instr), "instrText 里的 NBSP 必须原样保留"
    assert "尾巴 也有" in after(doc)[0], "域外那一个才该改"


def test_域内的显示文字不改(tmp_path, rules):
    """引用标注的显示文字（[1]）在域里，会被 Zotero 重新生成，改它没意义。"""
    path = make_doc(tmp_path, "正文")
    d = Document(path)
    add_field(d.paragraphs[0], display="[1]" + NBSP)
    d.save(path)

    doc, _, notes = run_fix(path, rules)
    assert NBSP in after(doc)[0]
    assert any("引用域" in n for n in notes)


# --------------------------------------------------------------------------
# 三、半角标点
# --------------------------------------------------------------------------
def test_括号成对全角化(tmp_path, rules):
    path = make_doc(tmp_path, "分布于南方省(区、市)一带")
    doc, changes, _ = run_fix(path, rules)
    assert after(doc)[0] == "分布于南方省（区、市）一带"
    assert any("全角化" in c.what for c in changes)


def test_括号跨节点也要成对改(tmp_path, rules):
    """左括号和右括号分在两个 w:t 里，也必须一起改。

    只改一边会变成"（区、市)"这种半角全角混排，比不改更难看。
    """
    path = make_doc(tmp_path, ["南方省(区、市", ")，其中"])
    doc, _, _ = run_fix(path, rules)
    assert after(doc)[0] == "南方省（区、市），其中"


def test_域内的标点整组跳过(tmp_path, rules):
    """括号落在域里（如域生成的括注）：整组放弃，不做半吊子修改。"""
    path = make_doc(tmp_path, "占位")
    d = Document(path)
    p = d.paragraphs[0]
    p.runs[0].text = "南方省"
    add_field(p, display="(区、市)")
    p.add_run("有分布")
    d.save(path)

    doc, changes, notes = run_fix(path, rules)
    assert after(doc)[0] == "南方省(区、市)有分布"
    assert not changes
    assert any("引用域" in n for n in notes)


def test_括号内含拉丁字母不改(tmp_path, rules):
    path = make_doc(tmp_path, "转录组测序(RNA sequencing)技术")
    doc, changes, _ = run_fix(path, rules)
    assert after(doc)[0] == "转录组测序(RNA sequencing)技术"
    assert not changes


def test_引用标注豁免(tmp_path, rules):
    path = make_doc(tmp_path, "已有报道[1,2]。")
    doc, changes, _ = run_fix(path, rules)
    assert after(doc)[0] == "已有报道[1,2]。"
    assert not changes


def test_普通冒号改_比例式冒号不改(tmp_path, rules):
    p1 = make_doc(tmp_path, "具有正负效应:一方面可诱导", name="a.docx")
    doc, _, _ = run_fix(p1, rules)
    assert after(doc)[0] == "具有正负效应：一方面可诱导"

    p2 = make_doc(tmp_path, "置于甲醛:乙酸:乙醇溶液中", name="b.docx")
    doc, _, _ = run_fix(p2, rules)
    assert after(doc)[0] == "置于甲醛:乙酸:乙醇溶液中", "比例式冒号留给人工判断"


def test_半角逗号分号全角化(tmp_path, rules):
    path = make_doc(tmp_path, "表现为正效应;另一方面,植物细胞失水")
    doc, _, _ = run_fix(path, rules)
    assert after(doc)[0] == "表现为正效应；另一方面，植物细胞失水"


def test_半角方括号降级为人工判断(tmp_path, rules):
    """全角方括号有 ［］ 与 【】 两种写法、语义不同，不机械替换。"""
    hi, lo, ex = punct_fixes("中文[中文]中文")
    assert not hi and lo, "方括号应落在人工判断那一堆"


# --------------------------------------------------------------------------
# 四、整层的不变量
# --------------------------------------------------------------------------
def test_run与w_t节点数量不变(tmp_path, rules):
    """只改文字，不新增、不合并、不删除 run。

    合并 run 会静默销毁 Zotero 引用域（实测 105 条 → 98 条，不报错）。
    """
    path = make_doc(tmp_path, ["图2.1", "油茶(果实)" + NBSP + "变化"])
    before = load(path)
    n_r = len(list(before.root.iter(f"{W}r")))
    n_t = len(list(before.root.iter(f"{W}t")))

    doc, _, _ = run_fix(path, rules)
    assert len(list(doc.root.iter(f"{W}r"))) == n_r
    assert len(list(doc.root.iter(f"{W}t"))) == n_t


def test_域标记与instrText逐字节不变(tmp_path, rules):
    path = make_doc(tmp_path, "图2.1油茶", "(国家基金)")
    d = Document(path)
    add_field(d.paragraphs[1])
    d.save(path)

    before = load(path)
    instr_before = [n.text for n in before.root.iter(f"{W}instrText")]
    fld_before = [n.get(f"{W}fldCharType") for n in before.root.iter(f"{W}fldChar")]

    doc, _, _ = run_fix(path, rules)
    assert [n.text for n in doc.root.iter(f"{W}instrText")] == instr_before
    assert [n.get(f"{W}fldCharType")
            for n in doc.root.iter(f"{W}fldChar")] == fld_before


def test_跑两遍没有新改动(tmp_path, rules):
    """幂等：把修完的文件再走一遍，文本层不该再报改动。

    注意不能在同一棵树上连跑两次——docmodel 的 p.text 是 load 时的快照，
    改完字不同步，第二次跑看到的还是旧文本，会重复插入（这个坑真踩到了）。
    所以必须写成文件、重新 load。
    """
    from formatter import apply

    src = make_doc(tmp_path, ["图2.1", "油茶(果实)" + NBSP + "变化"],
                   "进样量 10 " + MICRO_BAD + "L;见文献[1]。")
    first = tmp_path / "first.docx"
    apply(src, rules, first)

    changes2, _, _, _ = apply(first, rules, tmp_path / "second.docx",
                              dry_run=True)
    text2 = [c for c in changes2 if isinstance(c, TextChange)]
    assert not text2, f"第二次跑还报了文本改动：{[c.render() for c in text2]}"


# --------------------------------------------------------------------------
# 五、与检查器口径一致（防止两处判据各自演化）
# --------------------------------------------------------------------------
def test_检查器报的题注缺空格_修复后不再报(tmp_path, rules):
    from check_docx import run_check
    from formatter import apply

    path = make_doc(tmp_path, "图2.1油茶果实", "表3.1洗脱梯度", "图2.2 已有空格")
    rep, _, _ = run_check(str(path), rules)
    warned = [i for i in rep.items if i["category"].endswith("题格式")]
    assert len(warned) == 2, f"检查器应报 2 处，实际 {len(warned)}"

    out = tmp_path / "fixed.docx"
    apply(path, rules, out)
    rep2, _, _ = run_check(str(out), rules)
    assert not [i for i in rep2.items if i["category"].endswith("题格式")], \
        "报出来的那些必须全都被修掉"


def test_检查器报的异常字符_修复后不再报(tmp_path, rules):
    from check_docx import run_check
    from formatter import apply

    path = make_doc(tmp_path, "第一章" + PUA_JUNK + "绪论",
                    "进样量 10 " + MICRO_BAD + "L；另处 " + MICRO_GOOD + "m")
    rep, _, _ = run_check(str(path), rules)
    assert [i for i in rep.items if i["category"] == "异常字符"], "先得报出来"
    assert [i for i in rep.items if i["category"] == "字符混用"], "混用也得报"

    out = tmp_path / "fixed.docx"
    apply(path, rules, out)
    rep2, _, _ = run_check(str(out), rules)
    assert not [i for i in rep2.items if i["category"] == "异常字符"]
    assert not [i for i in rep2.items if i["category"] == "字符混用"]


def test_题注判据只有一份实现():
    """检查器与格式化器共用 caption_gap_pos，避免一处改了一处没改。"""
    assert caption_gap_pos("图2.1油茶", "图") is not None
    assert caption_gap_pos("图2.1 油茶", "图") is None
    assert caption_gap_pos("表3.1洗脱", "表") is not None
    assert caption_gap_pos("表3.1洗脱", "图") is None
    assert caption_gap_pos("图2.1", "图") is not None, "全程无空格也算缺"
