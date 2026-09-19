"""docmodel 的单元测试。

每条测试都对应一个真实踩过的坑——不是"测着玩"，是防止改坏。
命名用中文是为了可读，但不能带标点（Python 标识符限制）。
"""

from docx import Document
from docx.oxml.ns import qn

from conftest import build_docx, set_style_font
from docmodel import load


def test_段落总数含表格与内容控件(tmp_path):
    """python-docx 的 doc.paragraphs 只走 body 直接子级，
    表格和内容控件里的段落会整个漏掉。docmodel 必须全都收到。
    """
    doc = Document()
    doc.add_paragraph("正文一段")
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text = "单元格甲"
    t.cell(1, 1).text = "单元格乙"
    p = tmp_path / "t.docx"
    doc.save(p)

    d = load(p)
    body_only = len(Document(p).paragraphs)
    assert d.stats["paragraphs_in_table"] >= 4
    assert d.stats["paragraphs_total"] > body_only
    assert d.stats["paragraphs_body"] == body_only


def test_节数与sectPr元素数一致(tmp_path):
    """页边距检查曾用 doc.sections(13)，页码检查曾用 body.iter(sectPr)(14)，
    同一个报告里"节 2"不是同一个节。两个口径必须同源。
    """
    p = build_docx(tmp_path / "t.docx", paragraphs=[("第一章 绪论", "Heading 1")])
    d = load(p)
    assert d.stats["sections"] == len(d.sections)
    assert len(d.sections) >= 1


def test_页边距按twips换算而非EMU(tmp_path):
    """pgMar 的单位是 twips：1440 twips = 1 英寸 = 2.54cm。
    错用 EMU 换算会把 2.54cm 算成 0.004cm，页边距检查全线误报。
    """
    p = build_docx(tmp_path / "t.docx", margins_cm=(2.5, 2.5, 3.0, 2.5))
    m = load(p).sections[0].margins_cm
    assert abs(m["top"] - 2.5) < 0.02
    assert abs(m["bottom"] - 2.5) < 0.02
    assert abs(m["left"] - 3.0) < 0.02
    assert abs(m["right"] - 2.5) < 0.02


def test_页面尺寸按twips换算(tmp_path):
    p = build_docx(tmp_path / "t.docx", paragraphs=[("x",)])
    s = load(p).sections[0]
    assert abs(s.page_w_cm - 21.0) < 0.1
    assert abs(s.page_h_cm - 29.7) < 0.1


def test_无pStyle段落回退到默认样式(tmp_path):
    """没有 pStyle 的段落走的是文档默认样式（通常是 Normal）。
    不补这层回退，字号和行距会算成 None。
    """
    p = build_docx(tmp_path / "t.docx", paragraphs=[("一段普通文字",)])
    para = next(x for x in load(p).paragraphs if x.text.strip())
    assert para.size_pt is not None


def test_样式链走到头后回退到文档级默认(tmp_path):
    """字体/字号可能只写在 w:docDefaults 里。

    测出来的真 bug：原来的回退只到样式链为止，字号能出来但字体出不来。
    （主题字体是另一条路，见下一个测试。）
    """
    from conftest import set_doc_defaults

    p = build_docx(tmp_path / "t.docx", paragraphs=[("一段普通文字",)])
    set_doc_defaults(p, cn="仿宋", en="Arial", size_pt=14)

    para = next(x for x in load(p).paragraphs if x.text.strip())
    assert para.font_cn == "仿宋", "中文字体应从文档级默认取到"
    assert para.font_ascii == "Arial", "西文字体应从文档级默认取到"
    assert para.size_pt == 14.0


def test_主题字体引用能解析(tmp_path):
    """字体常以 w:asciiTheme="minorHAnsi" 引用，不是显式字体名。
    python-docx 自带模板就是这样，不解析就永远算作未设置。
    """
    from docx.oxml.ns import qn
    from docx import Document as D

    doc = D()
    para = doc.add_paragraph()
    r = para.add_run("text")
    rpr = r._element.get_or_add_rPr()
    rf = rpr.makeelement(qn("w:rFonts"), {})
    rf.set(qn("w:asciiTheme"), "minorHAnsi")
    rpr.insert(0, rf)
    p = tmp_path / "t.docx"
    doc.save(p)

    d = load(p)
    got = next(x for x in d.paragraphs if x.text.strip())
    assert got.font_ascii, "主题字体应被解析成实际字体名"


def test_中文字体与西文字体分别沿样式链回退(tmp_path):
    """某个样式可能只定义了 eastAsia 没定义 ascii——
    这时西文字体要继续往上找，不能直接判为 None。
    """
    doc = Document()
    set_style_font(doc, "Normal", cn="宋体", en="Times New Roman")
    set_style_font(doc, "Heading 1", cn="黑体")      # 故意不设西文
    doc.add_paragraph("第一章 绪论", style="Heading 1")
    p = tmp_path / "t.docx"
    doc.save(p)

    h = next(x for x in load(p).paragraphs if x.text.strip())
    assert h.font_cn == "黑体"
    assert h.font_ascii == "Times New Roman", "西文字体应沿样式链回退到 Normal"


def test_西文字体取样式定义里的值(tmp_path):
    """用户反馈的真实问题：二级标题的数字显示成宋体。
    根因是样式定义里 ascii=宋体，检查器必须能读出来。
    """
    doc = Document()
    set_style_font(doc, "Heading 2", cn="黑体", en="宋体")
    doc.add_paragraph("1.2 国内外研究现状", style="Heading 2")
    p = tmp_path / "t.docx"
    doc.save(p)

    h = next(x for x in load(p).paragraphs if x.text.strip())
    assert h.font_cn == "黑体"
    assert h.font_ascii == "宋体"


def test_引用域被标记(tmp_path):
    """含 Zotero 域的段落必须被标出来，否则格式化器可能把引用改坏。"""
    doc = Document()
    para = doc.add_paragraph()
    r = para.add_run()
    instr = r._element.makeelement(qn("w:instrText"), {})
    instr.text = " ADDIN ZOTERO_ITEM CSL_CITATION {} "
    r._element.append(instr)
    doc.add_paragraph("普通段落")
    p = tmp_path / "t.docx"
    doc.save(p)

    d = load(p)
    assert sum(1 for x in d.paragraphs if x.has_field) == 1
    assert len(d.fields) == 1


def test_样式名由styleId解析而来(tmp_path):
    """pStyle 存的是 id（常常是数字），不解析出来就看不到 "heading 1"、"toc 2"。"""
    doc = Document()
    doc.add_paragraph("第一章 绪论", style="Heading 1")
    p = tmp_path / "t.docx"
    doc.save(p)

    d = load(p)
    h = next(x for x in d.paragraphs if x.text.strip())
    assert h.style_id is not None
    assert (h.style_name or "").lower().startswith("heading")


def test_大纲级别继承自样式(tmp_path):
    """标题的大纲级别通常定义在样式里，段落本身没有。
    不沿样式链找，标题层级就整个判不出来（分类器全靠它）。
    """
    doc = Document()
    doc.add_paragraph("第一章 绪论", style="Heading 1")
    doc.add_paragraph("1.1 研究背景", style="Heading 2")
    p = tmp_path / "t.docx"
    doc.save(p)

    d = load(p)
    levels = {x.text: x.outline_level for x in d.paragraphs if x.text.strip()}
    assert levels["第一章 绪论"] == 0, "一级标题的大纲级别应为 0（来自样式）"
    assert levels["1.1 研究背景"] == 1, "二级标题的大纲级别应为 1（来自样式）"
