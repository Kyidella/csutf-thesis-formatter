"""检查器的集成测试：造一份"故意错"的文档，验证能不能抓到。

每条测试对应一个真实踩过或用户反馈过的问题。
"""

import pytest
import yaml

from check_docx import (FIXABLE_LABEL, FIX_KIND, MANUAL_LABEL, SKIP_FIX_KEYS,
                        classify_punct, cli_guard, fix_hint, fix_kind,
                        known_error, run_check)
from conftest import RULES, build_docx

pytestmark = pytest.mark.skipif(not RULES.exists(), reason="规则文件缺失")


@pytest.fixture(scope="module")
def rules():
    return yaml.safe_load(open(RULES, encoding="utf-8"))


def msgs(rep, category=None, level=None):
    return [i["msg"] for i in rep.items
            if (category is None or i["category"] == category)
            and (level is None or i["level"] == level)]


def test_页边距不符报错(tmp_path, rules):
    """Word 出厂默认是 2.54/3.17，规范要求 2.5/3.0。"""
    p = build_docx(tmp_path / "t.docx", margins_cm=(2.54, 2.54, 3.17, 3.17),
                   paragraphs=[("第一章 绪论", "Heading 1")])
    rep, _, _ = run_check(p, rules)
    assert msgs(rep, "页边距", "error"), "页边距不符必须报错"
    assert "2.54" in msgs(rep, "页边距", "error")[0]


def test_页边距正确不报错(tmp_path, rules):
    p = build_docx(tmp_path / "t.docx",
                   paragraphs=[("第一章 绪论", "Heading 1")])
    rep, _, _ = run_check(p, rules)
    assert not msgs(rep, "页边距"), "页边距合规时不该报任何东西"


def test_标题西文字体错误被抓到(tmp_path, rules):
    """用户反馈的真实问题：二级标题的数字显示成宋体。
    规范要求西文字体为 Times New Roman。
    """
    from docx import Document
    from conftest import set_style_font
    doc = Document()
    set_style_font(doc, "Heading 2", cn="黑体", en="宋体")
    doc.add_paragraph("1.2 国内外研究现状", style="Heading 2")
    p = tmp_path / "t.docx"
    doc.save(p)

    rep, _, _ = run_check(p, rules)
    hit = msgs(rep, "西文字体")
    assert hit, "二级标题西文字体为宋体必须被报出来"
    assert "宋体" in hit[0]


def test_标题西文字体正确不报错(tmp_path, rules):
    from docx import Document
    from conftest import set_style_font
    doc = Document()
    set_style_font(doc, "Heading 2", cn="黑体", en="Times New Roman")
    doc.add_paragraph("1.2 国内外研究现状", style="Heading 2")
    p = tmp_path / "t.docx"
    doc.save(p)

    rep, _, _ = run_check(p, rules)
    assert not msgs(rep, "西文字体")


def test_标题对齐不符被抓到(tmp_path, rules):
    """规范要求标题两端对齐，用左对齐会报出来。"""
    p = build_docx(tmp_path / "t.docx", paragraphs=[
        ("第一章 绪论", "Heading 1", "left", "黑体", "Times New Roman", 16)])
    rep, _, _ = run_check(p, rules)
    assert msgs(rep, "对齐")


def test_中英题注编号不一致报错(tmp_path, rules):
    p = build_docx(tmp_path / "t.docx", paragraphs=[
        ("图1.1 技术路线", None, "center", "宋体", "Times New Roman", 10.5),
        ("Fig 1.2 Technical route", None, "center", "宋体", "Times New Roman", 10.5),
    ])
    rep, _, _ = run_check(p, rules)
    assert msgs(rep, "中英题注不符", "error")


def test_图号后缺空格被抓到(tmp_path, rules):
    p = build_docx(tmp_path / "t.docx", paragraphs=[
        ("图1.1技术路线", None, "center", "宋体", "Times New Roman", 10.5)])
    rep, _, _ = run_check(p, rules)
    assert msgs(rep, "图题格式")


def test_缺失章节被抓到(tmp_path, rules):
    p = build_docx(tmp_path / "t.docx", paragraphs=[("随便正文。",)])
    rep, _, _ = run_check(p, rules)
    assert msgs(rep, "缺失章节")


def test_空段不计入格式统计(tmp_path, rules):
    """空行/占位段不承载格式意图。

    真 bug：空段被计入字号/字体统计后，会在报告里制造"规范要求12pt，
    实际2段24pt"这类假告警——而那 2 段根本是空行。
    """
    from docx.shared import Pt

    p = build_docx(tmp_path / "t.docx", paragraphs=[
        ("摘  要", None, "center", "黑体", "Times New Roman", 18),
        ("本文研究了油茶果实的滞育机制。", None, None, "宋体", "Times New Roman", 12),
    ])
    # 再手工塞一个"空段但带 24pt 直接格式"的段落
    from docx import Document
    doc = Document(p)
    empty = doc.add_paragraph(style="Normal")
    empty.add_run("").font.size = Pt(24)
    doc.save(p)

    rep, _, rows = run_check(p, rules)
    # 先确认这个空段确实被分到了摘要正文，否则测试没意义
    tail = [r for r in rows if r.para.text.strip() == ""]
    assert any(r.key == "abstract_body" for r in tail)

    # 摘要正文的字号检查不应出现（那条 24pt 是空段的，不算）
    body_size = [m for m in msgs(rep, "字号", "warn") if "中文摘要内容" in m]
    assert not body_size, f"空段的 24pt 不该计入摘要正文，实际：{body_size}"


# ---- 标点判定（纯函数） ----
def test_标点_两侧中文判为高置信度():
    hi, lo, ex = classify_punct("南方省(区、市)有大面积栽培")
    assert len(hi) == 1 and not lo


def test_标点_括号内含拉丁字母降级为人工判断():
    hi, lo, ex = classify_punct("测序(RNA sequencing)结果")
    assert not hi and len(lo) == 1


def test_标点_引用标注豁免():
    for text in ("见文献[1]。", "已有报道[8,9]。", "多项研究[20–22]。", "结果见[1,2,3]。"):
        hi, lo, ex = classify_punct(text)
        assert not hi and not lo, f"{text!r} 里的引用标注应豁免"
        assert ex == 1


def test_标点_英文环境不管():
    hi, lo, ex = classify_punct("Fig. 1 (a) shows the result.")
    assert not hi


def test_标点_比例式冒号降级():
    """甲醛:乙酸:乙醇 是比例式，半角冒号常见；而 正负效应:一方面 应改全角。"""
    hi, lo, ex = classify_punct("置于甲醛:乙酸:乙醇溶液中")
    assert not hi and len(lo) >= 1, "同段多个半角冒号判为比例式"

    hi2, lo2, ex2 = classify_punct("具有正负效应:一方面可诱导")
    assert len(hi2) == 1, "单个半角冒号应判为高置信度"


# ---- 异常字符（黑名单式，不能误伤正常字符） ----

def test_不换行空格被抓到(tmp_path, rules):
    """U+00A0 是 AI 生成文本的常见痕迹；WPS 里显示成 °。"""
    p = build_docx(tmp_path / "t.docx", paragraphs=[
        ("Camellia oleifera\u00a0Abel. is a woody plant.",)])
    rep, _, _ = run_check(p, rules)
    assert msgs(rep, "异常字符", "warn"), "NBSP 应被报出来"
    assert "不换行空格" in msgs(rep, "异常字符", "warn")[0]


def test_私有区字符被抓到(tmp_path, rules):
    """U+E5E5 是 WPS 留下的垃圾，在任何字体下都可能显示成空框。"""
    p = build_docx(tmp_path / "t.docx",
                   paragraphs=[("第一章\ue5e5绪论",)])
    rep, _, _ = run_check(p, rules)
    assert msgs(rep, "异常字符", "warn")


def test_同义字符混用被抓到(tmp_path, rules):
    """微符号(U+00B5) 与 希腊字母 μ(U+03BC) 都是"微"，同一篇里应统一。"""
    p = build_docx(tmp_path / "t.docx", paragraphs=[
        ("进样体积 10 \u00b5L，另一处写成 20 μL。",)])
    rep, _, _ = run_check(p, rules)
    assert msgs(rep, "字符混用", "warn")


def test_同义字符只用一种是干净的(tmp_path, rules):
    p = build_docx(tmp_path / "t.docx",
                   paragraphs=[("进样体积 10 μL。",)])
    rep, _, _ = run_check(p, rules)
    assert not msgs(rep, "字符混用")


def test_正常字符不误报(tmp_path, rules):
    """制表符、℃、外国人名重音字母、罗马数字、勾选框都是正常用法。

    用白名单报"一切非常见字符"会淹在这些里面——所以必须是黑名单。
    """
    p = build_docx(tmp_path / "t.docx", paragraphs=[
        ("\t\t目录项\t1",),
        ("10000 g、4℃离心 20 min。",),
        ("[78]\tALABADÍ D, BLÁZQUEZ M A. Molecular interaction.",),
        ("使用 HiScript Ⅲ QRT SuperMix。",),
        ("保密□，在      年解密后适用本授权书。",),
        ("μmol/g 与 °C 都正常。",),
    ])
    rep, _, _ = run_check(p, rules)
    assert not msgs(rep, "异常字符"), f"不该误报：{msgs(rep, '异常字符')}"
    assert not msgs(rep, "字符混用")


# ---- 报告里标出"这条能不能自动修" ----

def test_FIX_KIND_覆盖了所有检查分类():
    """新增一个检查项却忘了登记能不能自动修，这条会失败。

    从源码里抠出所有传给 rep.add / check_dist 的分类名，逐个查表。
    这段映射是被一次审查逼出来的：脚注编号检查器报 warn，格式化器里却根本没有
    写 footnotePr 的代码，光看报告看不出来。
    """
    import re
    from pathlib import Path

    src = (Path(__file__).resolve().parent.parent / "check_docx.py"
           ).read_text(encoding="utf-8")
    cats = set(re.findall(r'rep\.add\(\s*"[a-z]+",\s*"([^"]+)"', src))
    cats |= set(re.findall(r'check_dist\(\s*rep,\s*"([^"]+)"', src))
    # 下面这些是动态拼的或当参数传的，静态抠不出来，手工列全
    cats |= {"图题格式", "表题格式", "段前", "段后"}

    missing = cats - set(FIX_KIND)
    assert not missing, f"这些分类没登记能不能自动修：{sorted(missing)}"


def test_能自动修的标出层级_不能的标需人工():
    assert FIXABLE_LABEL in fix_hint("页边距")
    assert FIXABLE_LABEL in fix_hint("字号")
    assert FIXABLE_LABEL in fix_hint("图题格式")
    # 层次是内部概念，只出现在悬停提示里，不印在报告上
    assert fix_kind("字号") == "样式级"
    assert MANUAL_LABEL in fix_hint("脚注"), "脚注排版器确实不处理，必须标出来"
    assert MANUAL_LABEL in fix_hint("缺失章节")
    assert MANUAL_LABEL in fix_hint("某个还没实现的新分类"), "没登记的按保守处理"


def test_目录条目即便分类能修也要标需人工():
    """目录条目由 TOC 域生成，改了会被刷新覆盖，格式化器故意跳过。

    它是「字号」「对齐」这些能自动修的分类，只看分类会标成"支持自动排版"，
    用户点了自动排版却什么也没变——报告等于承诺了工具不会做的事。
    所以判定必须落到规则键上。这是拿真论文跑出来的：某样本论文排版完，
    剩下的警告几乎全是目录项，且全被标成了"支持自动排版"。
    """
    assert fix_kind("字号", "toc_entry_2") is None, "目录条目的字号修不了"
    assert MANUAL_LABEL in fix_hint("字号", "toc_entry_2")
    assert fix_kind("左缩进", "toc_entry_2") is None
    assert fix_kind("行距", "toc_entry_1") is None

    # 同一个分类落在正文标题上就是能修的——差别只在键
    assert fix_kind("字号", "heading2") == "样式级"
    assert FIXABLE_LABEL in fix_hint("字号", "heading2")

    # 少了键也不能反过来变宽松：目录条目的样式键都必须在跳过名单里
    assert {"toc_entry_front", "toc_entry_1", "toc_entry_2", "toc_entry_3"} \
        == SKIP_FIX_KEYS


def test_排版器跳过的目录条目名单和检查器是同一份():
    """两处各写一份名单，早晚会漂移——报告说能修、排版器说我不修。

    格式化器直接引用 check_docx 的常量，这条测试盯住这个引用别被改回去。
    """
    import formatter

    assert formatter.SKIP_FIX_KEYS is SKIP_FIX_KEYS


def test_命令行把认得的异常翻成中文_认不出的照旧traceback():
    """.doc 走命令行曾经直接甩 zipfile 的 traceback，界面里却是人话。

    现在两边共用一张表：认得的翻成中文退出；认不出的原样抛——那才是真出了
    bug，把 traceback 藏掉反而没法排查。
    """
    import zipfile

    def 打开一个doc():
        raise zipfile.BadZipFile()

    with pytest.raises(SystemExit) as e:
        cli_guard(打开一个doc)
    assert ".docx" in str(e.value), "要告诉用户这可能是 .doc"
    assert "Traceback" not in str(e.value)

    def 真出了bug():
        raise RuntimeError("没见过这种")

    with pytest.raises(RuntimeError):
        cli_guard(真出了bug)          # 不认得的异常不该被吞掉

    assert known_error(zipfile.BadZipFile())
    assert known_error(FileNotFoundError())
    assert not known_error(RuntimeError("没见过这种"))


def test_报告里带上了能不能自动修的标记(tmp_path, rules):
    """标记要真的出现在写出的报告里，不能只是个函数。"""
    from check_docx import write_report

    p = build_docx(tmp_path / "t.docx", margins_cm=(2.5, 2.5, 2.5, 2.5))
    out = tmp_path / "r.md"
    rep, doc, rows = run_check(p, rules, str(out))
    text = out.read_text(encoding="utf-8")
    assert FIXABLE_LABEL in text and MANUAL_LABEL in text, "报告里两种标记都该出现"
    assert not any("{'" in line for line in text.splitlines()), \
        "报告里不该再出现原始字典（规则路径那个 bug 修过了）"


def test_进度回调按顺序推阶段(tmp_path, rules):
    """GUI 靠这个回调推进度条，断了界面就一直卡在 0。

    约定：分母固定，分子从 0 走到总数；没传回调时行为完全不变。
    """
    p = build_docx(tmp_path / "t.docx", paragraphs=[("正文内容。",)])
    seen = []
    run_check(p, rules, on_stage=lambda msg, done, total: seen.append((done, total, msg)))

    assert seen, "一个阶段都没推"
    assert seen[0][0] == 0, "第一次应该是 0"
    assert seen[-1][0] == seen[-1][1], "最后一次应该走满"
    assert len({t for _, t, _ in seen}) == 1, "分母中途变了"
    assert all(isinstance(m, str) and m for _, _, m in seen)


def test_不传回调也能正常跑(tmp_path, rules):
    p = build_docx(tmp_path / "t.docx", paragraphs=[("正文内容。",)])
    rep, doc, rows = run_check(p, rules)
    assert rep.items is not None and rows


def test_图片段落不计入样式统计(tmp_path, rules):
    """图片段落排版时有意跳过（套正文样式会把嵌入型图片裁掉），
    那就不该再拿样式规范去要求它——报了用户也没法改，只是噪声。

    真踩过：跳过图片段落之后，"图题行距不符"多出来一条，
    指的是一个题注和图片挤在同一段的段落。
    """
    from docx import Document
    from docx.shared import Pt
    from conftest import add_fake_image

    src = build_docx(tmp_path / "t.docx", paragraphs=[
        ("图1.1 技术路线", None, "center", "宋体", "Times New Roman", 10.5),
    ])
    d = Document(src)
    p = d.add_paragraph("图1.2 带图的题注")
    p.paragraph_format.line_spacing = 1.5          # 故意不合规范
    p.paragraph_format.first_line_indent = Pt(24)
    add_fake_image(p)
    d.save(src)

    rep, _, _ = run_check(src, rules)
    bad = [i for i in rep.items
           if i["category"] in ("行距", "首行缩进") and "图题" in i["msg"]]
    assert not bad, f"图片段落不该被算进样式统计：{bad}"

    # 而没有图片的那个题注仍然是合规的，不该因此漏报别的
    assert not [i for i in rep.items if i["category"] == "字号" and "图题" in i["msg"]]


def test_标题和正文挤在同一段会被报出来(tmp_path, rules):
    """软回车把标题和正文连成一段，分类器只能当正文——标题样式套不上。

    真踩过：用户看到"3.2.3 没排上"，以为工具漏了；其实是那份稿子里
    标题和正文在同一个段落里（Word 里按的是 Shift+Enter）。
    """
    from docx import Document
    p = build_docx(tmp_path / "t.docx", paragraphs=[("正文一。",)])
    d = Document(p)
    par = d.add_paragraph()
    par.add_run("3.2.3 镉胁迫触发木质素合成基因的转录上调")
    par.add_run().add_break()                       # 软回车
    par.add_run("为进一步了解镉耐受的分子机制，我们分析了相关基因的表达谱。")
    d.save(p)

    rep, _, _ = run_check(p, rules)
    hit = msgs(rep, "标题与正文同段", "warn")
    assert hit, "标题和正文同段必须报出来"
    assert "软回车" in hit[0]


def test_正常分段不会被误报(tmp_path, rules):
    """标题自己一段、正文自己一段——这是正常的，不该报。"""
    p = build_docx(tmp_path / "t.docx", paragraphs=[
        ("3.2.3 氧化应激响应",),
        ("为进一步了解镉耐受的分子机制。",),
    ])
    rep, _, _ = run_check(p, rules)
    assert not msgs(rep, "标题与正文同段")


def test_段内有换行但首行不是标题则不该报(tmp_path, rules):
    """正文里用软回车手动换行很常见，首行不是编号标题就不是"标题与正文同段"。

    （这条是变异测试逼出来的：判定条件写坏成"恒真"时，上一版测试抓不到。）
    """
    from docx import Document
    p = build_docx(tmp_path / "t.docx", paragraphs=[("正文一。",)])
    d = Document(p)
    par = d.add_paragraph()
    par.add_run("这一行是正文的继续，不是标题")
    par.add_run().add_break()
    par.add_run("换行之后的又一段文字。")
    d.save(p)

    rep, _, _ = run_check(p, rules)
    assert not msgs(rep, "标题与正文同段"), \
        f"普通换行不该被当成标题与正文同段：{msgs(rep, '标题与正文同段')}"
