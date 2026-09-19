"""文档级格式化器的测试。

这一层动的是 sectPr，不碰正文 run。测试要盯住三件事：
  1. 该改的改了
  2. 不该动的字节级没动（其余 zip 部件哈希不变）
  3. 引用域完好
"""

import hashlib
import zipfile

import pytest
import yaml
from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from conftest import RULES, build_docx
from formatter import apply
from docmodel import TWIPS_PER_CM

pytestmark = pytest.mark.skipif(not RULES.exists(), reason="规则文件缺失")


@pytest.fixture(scope="module")
def rules():
    return yaml.safe_load(open(RULES, encoding="utf-8"))


def twips_to_cm(t):
    return round(int(t) / TWIPS_PER_CM, 2)


def part_hashes(path, exclude=("word/document.xml", "word/styles.xml")):
    """除被替换的那两个部件外，其余部件的哈希。"""
    z = zipfile.ZipFile(path)
    return {n: hashlib.sha256(z.read(n)).hexdigest()
            for n in z.namelist() if n not in exclude}


def add_field_paragraph(doc):
    """塞一个 Zotero 引用域，用来验证格式化不动它。"""
    p = doc.add_paragraph()
    r = p.add_run()
    instr = r._element.makeelement(qn("w:instrText"), {})
    instr.text = " ADDIN ZOTERO_ITEM CSL_CITATION {} "
    r._element.append(instr)


# --------------------------------------------------------------------------
def test_页边距被改成规范值(tmp_path, rules):
    src = build_docx(tmp_path / "a.docx", margins_cm=(2.54, 2.54, 3.17, 3.17),
                     paragraphs=[("第一章 绪论", "Heading 1")])
    dst = tmp_path / "b.docx"
    changes, _, _, _ = apply(src, rules, dst)
    assert any("边距" in c.what for c in changes)

    from docmodel import load
    m = load(dst).sections[0].margins_cm
    assert abs(m["top"] - 2.5) < 0.02
    assert abs(m["left"] - 3.0) < 0.02
    assert abs(m["right"] - 2.5) < 0.02


def test_页面尺寸改成A4(tmp_path, rules):
    """Letter 纸必须被改成 A4。

    注意 fixture 要显式给 Letter，否则源文件本来就是 A4，
    "宽高写反"这类 bug 会被漏掉（这正是一次漏网发现的）。
    """
    src = build_docx(tmp_path / "a.docx", page_cm=(21.59, 27.94),
                     paragraphs=[("正文",)])
    dst = tmp_path / "b.docx"
    changes, _, _, _ = apply(src, rules, dst)
    assert any("页面尺寸" in c.what for c in changes), "应报出尺寸改动"

    from docmodel import load
    s = load(dst).sections[0]
    assert abs(s.page_w_cm - 21.0) < 0.05, f"宽应为 21.0cm，实际 {s.page_w_cm}"
    assert abs(s.page_h_cm - 29.7) < 0.05, f"高应为 29.7cm，实际 {s.page_h_cm}"


def test_数值等价的边距不算改动(tmp_path, rules):
    """页脚边距 851 twips 与 850 twips 都是 1.5cm，不该被报成"1.5cm → 1.5cm"。

    真踩过：按字符串比较会报出这种假改动。
    """
    from conftest import set_sect_twips

    src = build_docx(tmp_path / "a.docx", paragraphs=[("正文",)])
    set_sect_twips(src, footer=851, header=850,
                   top=1417, bottom=1417, left=1701, right=1417)
    dst = tmp_path / "b.docx"
    changes, _, _, _ = apply(src, rules, dst)
    spurious = [c.render() for c in changes
                if c.old == c.new or c.old.replace("cm", "") ==
                c.new.replace("cm", "")]
    assert not spurious, f"不该出现新旧相同的假改动：{spurious}"


def test_页码按前置与正文分段(tmp_path, rules):
    """规范：摘要到目录用大写罗马数字，正文起阿拉伯数字重新从 1 开始。"""
    doc = Document()
    doc.add_paragraph("封面")
    for text in ("摘  要", "目  录", "第一章 绪论"):
        doc.add_section(WD_SECTION.NEW_PAGE)
        doc.add_paragraph(text)
    src = tmp_path / "a.docx"
    doc.save(src)
    dst = tmp_path / "b.docx"

    apply(src, rules, dst)
    from docmodel import load
    fmts = {s.index: s.pg_num_fmt for s in load(dst).sections}
    assert fmts[1] == "upperRoman", "摘要所在节应为大写罗马数字"
    assert fmts[2] == "upperRoman", "目录所在节应为大写罗马数字"
    assert fmts[3] == "decimal", "正文所在节应为阿拉伯数字"
    assert fmts[0] is None or fmts[0] != "upperRoman", "封面不动"


def test_只改两个xml其余部件逐字节不变(tmp_path, rules):
    """这是"没动到别的东西"的硬证据。

    格式化器只替换 word/document.xml（节与页码）和 word/styles.xml（样式定义），
    图片、字体、批注、页眉页脚部件、settings.xml 等一律逐字节照抄。
    """
    src = build_docx(tmp_path / "a.docx", paragraphs=[("第一章 绪论", "Heading 1")])
    dst = tmp_path / "b.docx"
    before = part_hashes(src)
    apply(src, rules, dst)
    after = part_hashes(dst)
    assert before.keys() == after.keys(), "部件数量不能变"
    changed = [n for n in before if before[n] != after[n]]
    assert changed == [], f"这些部件不该被改动：{changed}"


def test_引用域完好(tmp_path, rules):
    doc = Document()
    add_field_paragraph(doc)
    doc.add_paragraph("第一章 绪论")
    src = tmp_path / "a.docx"
    doc.save(src)
    dst = tmp_path / "b.docx"

    def counts(p):
        x = zipfile.ZipFile(p).read("word/document.xml").decode("utf-8")
        return (x.count("ZOTERO_ITEM"), x.count('w:fldCharType="begin"'),
                x.count('w:fldCharType="end"'))

    before = counts(src)
    apply(src, rules, dst)
    assert counts(dst) == before, "引用域结构不能变"


def test_拒绝原地修改(tmp_path, rules):
    src = build_docx(tmp_path / "a.docx", paragraphs=[("正文",)])
    with pytest.raises(ValueError):
        apply(src, rules, src)


def test_预演模式不写文件(tmp_path, rules):
    src = build_docx(tmp_path / "a.docx", paragraphs=[("正文",)])
    dst = tmp_path / "b.docx"
    changes, _, _, _ = apply(src, rules, dst, dry_run=True)
    assert changes, "预演也应报告会改什么"
    assert not dst.exists(), "预演不应写出文件"


def test_已合规的文件不产生改动(tmp_path, rules):
    """幂等性：跑过一遍之后再跑，应该没有改动。"""
    src = build_docx(tmp_path / "a.docx", paragraphs=[("第一章 绪论", "Heading 1")])
    once = tmp_path / "once.docx"
    twice = tmp_path / "twice.docx"

    apply(src, rules, once)
    changes, _, _, _ = apply(once, rules, twice)
    assert changes == [], f"第二次不该再有改动，实际：{[c.render() for c in changes]}"


def test_横向页不改尺寸(tmp_path, rules):
    """横向页在论文里通常是刻意的（大幅插图），不该被拉回 A4 纵向。"""
    from docx.enum.section import WD_ORIENT
    doc = Document()
    doc.add_paragraph("横向插图页")
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    src = tmp_path / "a.docx"
    doc.save(src)
    dst = tmp_path / "b.docx"

    changes, notes, _, _ = apply(src, rules, dst)
    assert not any("页面尺寸" in c.what for c in changes)
    assert any("横向" in n for n in notes)

    from docmodel import load
    s = load(dst).sections[0]
    assert abs(s.page_w_cm - 29.7) < 0.05, "横向页尺寸不该被改"


# --------------------------------------------------------------------------
# 段落映射（4b）
# --------------------------------------------------------------------------
def pstyles(path):
    """段落序号 → 该段的 pStyle 值。"""
    from lxml import etree
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    root = etree.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    out = []
    for p in root.iter(f"{W}p"):
        st = p.find(f"{W}pPr/{W}pStyle")
        out.append(st.get(f"{W}val") if st is not None else None)
    return out


def test_段落被指向规范样式(tmp_path, rules):
    src = build_docx(tmp_path / "a.docx", paragraphs=[
        ("摘  要", None, "center", "黑体", "Times New Roman", 18),
        ("本文研究了油茶果实的滞育机制。", None, None, "宋体", "Times New Roman", 12),
    ])
    dst = tmp_path / "b.docx"
    changes, _, style_map, stat = apply(src, rules, dst)

    assert stat["styled"] > 0, "应有段落被指向样式"
    styles = pstyles(dst)
    assert style_map["abstract_title"] in styles
    assert style_map["abstract_body"] in styles


def test_含引用域的段落被跳过(tmp_path, rules):
    """格式改动会被 Zotero 刷新覆盖，且会触发它的提示，所以跳过。"""
    doc = Document()
    doc.add_paragraph("摘  要", style="Heading 1")
    doc.add_paragraph("摘要正文。")
    add_field_paragraph(doc)          # 含 ZOTERO_ITEM 域
    src = tmp_path / "a.docx"
    doc.save(src)
    dst = tmp_path / "b.docx"

    _, notes, _, stat = apply(src, rules, dst)
    assert stat["skip_field"] == 1
    assert any("引用域" in n for n in notes)


def test_目录条目被跳过(tmp_path, rules):
    """目录条目由 TOC 域生成，改 pStyle 会在刷新目录时被覆盖。

    真实目录条目带点前导符（或使用 toc 1/2/3 样式）；这里用点前导符，
    否则「第一章 绪论」会被当成章标题。
    """
    doc = Document()
    doc.add_paragraph("目  录")
    doc.add_paragraph("第一章 绪论................1")
    src = tmp_path / "a.docx"
    doc.save(src)
    dst = tmp_path / "b.docx"

    _, _, _, stat = apply(src, rules, dst)
    assert stat["skip_toc"] >= 1, "带点前导符的目录条目应被判为目录区"


def test_冲突的直接格式被清除(tmp_path, rules):
    """段落上若留着错误的直接字号，样式就算写对了也盖不住。"""
    src = build_docx(tmp_path / "a.docx", paragraphs=[
        ("第一章 绪论", "Heading 1", "left", "宋体", "宋体", 24),
    ])
    dst = tmp_path / "b.docx"
    apply(src, rules, dst)

    from docmodel import load
    from classifier import classify
    d = load(dst)
    h = next(p for p in d.paragraphs if p.text.strip())
    # 直接格式被清掉后，有效格式应回到样式定义的 黑体/16pt
    assert h.size_pt == 16.0, f"字号应回到样式的 16pt，实际 {h.size_pt}"
    assert h.font_cn == "黑体", f"中文字体应回到样式的黑体，实际 {h.font_cn}"


def test_斜体等语义格式不被清除(tmp_path, rules):
    """只清我们管的属性；加粗、斜体、上下标有语义，必须留着。"""
    doc = Document()
    p = doc.add_paragraph("本文研究了油茶果实（*Camellia oleifera*）的滞育。")
    for r in p.runs:
        pass
    r = p.add_run("斜体术语")
    r.italic = True
    r.font.size = Pt(10.5)            # 故意给个错的字号
    src = tmp_path / "a.docx"
    doc.save(src)
    dst = tmp_path / "b.docx"
    apply(src, rules, dst)

    from docmodel import load
    d = load(dst)
    italics = [run for para in d.paragraphs for run in para.element.findall(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}r")
        if run.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
                    "rPr/{http://schemas.openxmlformats.org/wordprocessingml/2006/main}i")
        is not None]
    assert italics, "斜体必须保留"



def test_styles_xml仍是合法样式表(tmp_path, rules):
    """写回时若把 document.xml 的内容错写到 styles.xml，这里必须炸。

    两个 xml 都被排除在"其余部件哈希不变"的比对之外，
    所以需要单独断言 styles.xml 本身还是合法的样式表。
    """
    from lxml import etree
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

    src = build_docx(tmp_path / "a.docx", paragraphs=[("第一章 绪论", "Heading 1")])
    dst = tmp_path / "b.docx"
    apply(src, rules, dst)

    data = zipfile.ZipFile(dst).read("word/styles.xml")
    root = etree.fromstring(data)                 # 坏了会在这里抛
    assert etree.QName(root).localname == "styles"
    assert root.findall(f"{W}style"), "styles.xml 里应还有样式定义"

    doc = etree.fromstring(zipfile.ZipFile(dst).read("word/document.xml"))
    assert etree.QName(doc).localname == "document", "两个 xml 不能写串"


def test_格式化也有进度回调(tmp_path, rules):
    src = build_docx(tmp_path / "a.docx", paragraphs=[("第一章 绪论", "Heading 1")])
    seen = []
    apply(src, rules, tmp_path / "b.docx", dry_run=True,
          on_stage=lambda msg, done, total: seen.append((done, total, msg)))
    assert seen and seen[0][0] == 0
    assert all(t == seen[0][1] for _, t, _ in seen), "分母中途变了"
    assert any("文本" in m for _, _, m in seen), "应该有一个阶段是修文本"


# ---- 图片段落：不能套样式，固定行距要救回来 ----

from conftest import add_fake_image as add_inline_image


def test_图片段落不套正文样式(tmp_path, rules):
    """正文样式带着"固定行距 20 磅"，会把嵌入型图片裁掉。

    实测：某样本论文 31 张内嵌图在完整排版输出里全被套上了固定行距。
    这里是回归测试——图片段落必须原样不动。
    """
    src = build_docx(tmp_path / "a.docx", paragraphs=[
        ("摘  要", None, "center", "黑体", "Times New Roman", 18),
        ("正文段落。", None, None, "宋体", "Times New Roman", 12),
    ])
    from docx import Document
    d = Document(src)
    img = d.add_paragraph()
    add_inline_image(img)
    d.save(src)

    dst = tmp_path / "b.docx"
    apply(src, rules, dst)

    from docmodel import load
    doc = load(dst)
    for p in doc.paragraphs:
        if any(n.tag.split("}")[-1] in ("drawing", "blip")
               for n in p.element.iter()):
            ls = p.pformat.get("line_spacing")
            assert not (ls and ls[0] == "exact"), \
                f"段{p.index} 是图片段落，却被套上了固定行距 {ls}——图会被裁掉"
            assert p.style_id is None, f"段{p.index} 图片段落不该被指派样式"


def test_图片段落上的固定行距会被改回单倍(tmp_path, rules):
    """已经带着固定行距的图片段落（比如上一版工具处理过的文件）要救回来。"""
    src = build_docx(tmp_path / "a.docx", paragraphs=[
        ("摘  要", None, "center", "黑体", "Times New Roman", 18),
    ])
    from docx import Document
    from docx.shared import Pt
    d = Document(src)
    img = d.add_paragraph()
    img.paragraph_format.line_spacing = Pt(20)      # 固定行距
    img.paragraph_format.line_spacing_rule = 4      # EXACTLY
    add_inline_image(img)
    d.save(src)

    dst = tmp_path / "b.docx"
    changes, notes, _, stat = apply(src, rules, dst)
    assert stat["unclipped"] == 1, "该报出救回了一段"

    from docmodel import load
    ls = load(dst).paragraphs[1].pformat.get("line_spacing")
    assert not (ls and ls[0] == "exact"), f"固定行距没清掉：{ls}"


def test_缺内容会在说明里写出来(tmp_path, rules):
    """在缺摘要、缺参考文献的稿子上排版，改动清单看着一切正常，
    用户会以为文档没问题。必须明说缺什么。

    真踩过：用户拿一份缺了大半章节的稿子测试，界面上只有一份正常的改动清单。
    """
    src = build_docx(tmp_path / "a.docx", paragraphs=[("随便一段正文。",)])
    dst = tmp_path / "b.docx"
    changes, notes, _, stat = apply(src, rules, dst)

    miss = stat.get("missing_sections")
    assert miss, "缺章节没被记进 stat，界面就没法醒目提示"
    assert "摘要" in miss and "参考文献" in miss
    assert any("没有检测到" in n for n in notes), f"说明里没写：{notes}"


def test_有的章节不会被报成缺失(tmp_path, rules):
    """备齐的那几章不该出现在缺失列表里（附录、ABSTRACT 没给，本来就该报）。"""
    src = build_docx(tmp_path / "a.docx", paragraphs=[
        ("摘  要", None, "center", "黑体", "Times New Roman", 18),
        ("中文摘要内容。", None, None, "宋体", "Times New Roman", 12),
        ("目  录", None, "center", "黑体", "Times New Roman", 16),
        ("参考文献", None, "center", "黑体", "Times New Roman", 16),
        ("[1] 某某. 文章[J]. 期刊, 2024.", None, None, "宋体", "Times New Roman", 12),
        ("致  谢", None, "center", "黑体", "Times New Roman", 16),
        ("感谢。", None, None, "宋体", "Times New Roman", 12),
    ])
    _, notes, _, stat = apply(src, rules, tmp_path / "b.docx")
    for have in ("摘要", "目录", "参考文献", "致谢"):
        assert have not in stat["missing_sections"], f"{have} 明明有，却被报成缺失"
    # 没给的那些仍然要报，而且说明里要点名
    assert "附录A" in stat["missing_sections"]
    assert any("附录A" in n for n in notes)


def test_图片和文字同段时仍然排版_但行距被钉成单倍(tmp_path, rules):
    """图片锚在标题/题注里时，整段跳过会把标题也跳过（用户实测"3.3 没排上"）。

    改成：有文字的照常排版，只把行距钉成单倍防裁图。
    """
    from docx import Document
    from conftest import add_fake_image

    src = build_docx(tmp_path / "a.docx", paragraphs=[
        ("摘  要", None, "center", "黑体", "Times New Roman", 18),
        ("图2.1 带图的题注", None, "center", "宋体", "Times New Roman", 10.5),
    ])
    d = Document(src)
    add_fake_image(d.paragraphs[1])          # 题注和嵌入图挤在一段
    d.save(src)

    dst = tmp_path / "b.docx"
    changes, notes, _, stat = apply(src, rules, dst)
    assert stat["skip_image"] == 0, "有文字的图片段落不该被跳过"
    assert stat["unclipped"] == 1, "该给它钉一个单倍行距"

    from docmodel import load, W
    p = load(dst).paragraphs[1]
    assert p.style_id, "有文字的图片段落应该照常套样式"
    # docmodel 把 lineRule=auto 换算成倍数，所以"单倍"在这里是 ("multiple", 1.0)
    ls = p.pformat.get("line_spacing")
    assert ls == ("multiple", 1.0),         f"行距该被钉成单倍，实际 {ls}（固定行距会裁掉嵌入图）"


def _style_bold_in(dst, style_name):
    """从写出的 styles.xml 里读某个样式的加粗开关。"""
    import zipfile
    from lxml import etree
    from docmodel import W
    root = etree.fromstring(zipfile.ZipFile(dst).read("word/styles.xml"))
    for s in root.findall(f"{W}style"):
        nm = s.find(f"{W}name")
        if nm is not None and (nm.get(f"{W}val") or "").lower() == style_name.lower():
            b = s.find(f"{W}rPr/{W}b")
            if b is None:
                return None
            return b.get(f"{W}val") not in ("0", "false")
    raise AssertionError(f"没找到样式 {style_name}")


def test_标题样式被显式关掉加粗(tmp_path, rules):
    """Word 内置的 heading 2/3/4 自带 <w:b/>。

    我们改了字体字号却没管加粗，结果标题被我们弄成了粗体——规范要的是
    "黑体 三号"，从没说加粗（黑体本身已经够重）。用户实测发现的。
    """
    src = build_docx(tmp_path / "a.docx", paragraphs=[
        ("1.2 国内外研究现状", "Heading 2"),
        ("1.2.1 小节标题", "Heading 3"),
    ])
    dst = tmp_path / "b.docx"
    apply(src, rules, dst)
    for name in ("heading 2", "heading 3"):
        assert _style_bold_in(dst, name) is False, \
            f"{name} 必须显式关掉加粗（不写就会继承内置样式的加粗）"


def test_ABSTRACT标题按规范保持加粗(tmp_path, rules):
    """规范明确要求 ABSTRACT 加粗——别把"标题都不加粗"一刀切。
    （这条原先登记在 not_implemented 里，现在实现了。）"""
    src = build_docx(tmp_path / "a.docx", paragraphs=[
        ("ABSTRACT", None, "center", None, "Times New Roman", 18),
    ])
    dst = tmp_path / "b.docx"
    apply(src, rules, dst)
    assert _style_bold_in(dst, "ABSTRACT 标题") is True
