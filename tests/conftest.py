"""测试公共工具：合成 .docx 与手工构造 Para。

合成文档用 python-docx 建，因为要精确控制"错的格式"来验证检查器能不能抓到。
"""

import sys
from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

SAMPLES = {
    # 真实样本不入库（含姓名与未发表内容）。要跑这些回归测试，把文件按下面的
    # 名字放到 ROOT.parent / "论文" 下即可；不在就自动跳过。
    "终版": ROOT.parent / "论文" / "样本论文.docx",
    "送审版": ROOT.parent / "论文" / "样本论文（送审稿）.docx",
}
RULES = ROOT / "rules" / "中南林科大_硕士_学术学位_理工类.yaml"


def set_font(run, cn=None, en=None, size_pt=None, bold=None):
    """同时设置中文字体(eastAsia)与西文字体(ascii/hAnsi)。

    python-docx 的 run.font.name 只写 ascii/hAnsi，设不了 eastAsia，
    所以这里直接操作 XML。
    """
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rf)
    if cn:
        rf.set(qn("w:eastAsia"), cn)
    if en:
        rf.set(qn("w:ascii"), en)
        rf.set(qn("w:hAnsi"), en)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if bold is not None:
        run.font.bold = bold


def set_style_font(doc, style_name, cn=None, en=None, size_pt=None):
    """改样式定义里的字体——这是最容易被漏掉的一层。"""
    st = doc.styles[style_name]
    rpr = st.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = rpr.makeelement(qn("w:rFonts"), {})
        rpr.insert(0, rf)
    if cn:
        rf.set(qn("w:eastAsia"), cn)
    if en:
        rf.set(qn("w:ascii"), en)
        rf.set(qn("w:hAnsi"), en)
    if size_pt is not None:
        st.font.size = Pt(size_pt)


def build_docx(path, *, margins_cm=(2.5, 2.5, 3.0, 2.5), page_cm=(21.0, 29.7),
               paragraphs=(), add_table=False):
    """按配方造一个 .docx。

    paragraphs 每项是 (text, style, align, cn, en, size_pt)，
    后面几项可省略，例如 ("第一章 绪论", "Heading 1")。
    """
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    top, bottom, left, right = margins_cm
    for sec in doc.sections:
        sec.top_margin, sec.bottom_margin = Cm(top), Cm(bottom)
        sec.left_margin, sec.right_margin = Cm(left), Cm(right)
        # python-docx 默认是 Letter 纸；page_cm 可覆盖，用来测"非 A4 应被改"
        sec.page_width, sec.page_height = Cm(page_cm[0]), Cm(page_cm[1])

    ALIGN = {"center": WD_ALIGN_PARAGRAPH.CENTER,
             "left": WD_ALIGN_PARAGRAPH.LEFT,
             "both": WD_ALIGN_PARAGRAPH.JUSTIFY}

    for spec in paragraphs:
        text, style, align, cn, en, size = (list(spec) + [None] * 6)[:6]
        p = doc.add_paragraph(text, style=style) if style else doc.add_paragraph(text)
        if align:
            p.alignment = ALIGN[align]
        for r in p.runs:
            set_font(r, cn=cn, en=en, size_pt=size)

    if add_table:
        t = doc.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "1.5"
        t.cell(0, 1).text = "1.2"

    doc.save(path)
    return path


def set_doc_defaults(path, cn=None, en=None, size_pt=None):
    """改 .docx 里的文档级默认格式（w:docDefaults），直接改 zip 里的 styles.xml。

    这一层是字体/字号的最后兜底：样式链走到头之后就看它。
    """
    import shutil
    import zipfile
    from lxml import etree

    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    src = Path(path).with_suffix(".orig.docx")
    shutil.move(path, src)
    zin = zipfile.ZipFile(src)
    root = etree.fromstring(zin.read("word/styles.xml"))

    rpr = root.find(f"{W}docDefaults/{W}rPrDefault/{W}rPr")
    if rpr is None:
        rpr = etree.SubElement(
            etree.SubElement(etree.SubElement(root, f"{W}docDefaults"),
                             f"{W}rPrDefault"), f"{W}rPr")
    rf = rpr.find(f"{W}rFonts")
    if rf is None:
        rf = rpr.makeelement(f"{W}rFonts", {})
        rpr.insert(0, rf)
    if cn:
        rf.set(f"{W}eastAsia", cn)
    if en:
        rf.set(f"{W}ascii", en)
        rf.set(f"{W}hAnsi", en)
    if size_pt is not None:
        sz = rpr.find(f"{W}sz")
        if sz is None:
            sz = rpr.makeelement(f"{W}sz", {})
            rpr.append(sz)
        sz.set(f"{W}val", str(int(size_pt * 2)))

    new_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                             standalone=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = new_xml if item.filename == "word/styles.xml" \
                else zin.read(item.filename)
            zout.writestr(item, data)
    zin.close()          # Windows 上不关就删不掉源文件
    src.unlink()
    return path


def set_sect_twips(path, **attrs):
    """直接往 sectPr 的 pgMar / pgSz 上写指定的 twips 值。

    用来构造"数值上等价但写法不同"的情形，验证比较是否有容差。
    """
    import shutil
    import zipfile
    from lxml import etree

    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    src = Path(path).with_suffix(".orig.docx")
    shutil.move(path, src)
    zin = zipfile.ZipFile(src)
    root = etree.fromstring(zin.read("word/document.xml"))

    for sect in root.iter(f"{W}sectPr"):
        for key, val in attrs.items():
            if key in ("w", "h"):
                node = sect.find(f"{W}pgSz")
                if node is None:
                    node = sect.makeelement(f"{W}pgSz", {})
                    sect.insert(0, node)
            else:
                node = sect.find(f"{W}pgMar")
                if node is None:
                    node = sect.makeelement(f"{W}pgMar", {})
                    sect.append(node)
            node.set(f"{W}{key}", str(val))

    new_xml = etree.tostring(root, xml_declaration=True, encoding="UTF-8",
                             standalone=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = (new_xml if item.filename == "word/document.xml"
                    else zin.read(item.filename))
            zout.writestr(item, data)
    zin.close()
    src.unlink()
    return path


@pytest.fixture
def tmp_docx(tmp_path):
    def _make(name="test.docx", **kw):
        return build_docx(tmp_path / name, **kw)
    return _make


@pytest.fixture(params=list(SAMPLES))
def sample(request):
    """真实论文样本；文件不在就跳过（隐私原因样本不入库）。"""
    p = SAMPLES[request.param]
    if not p.exists():
        pytest.skip(f"样本不存在：{p}")
    return request.param, p


def make_para(index, text, **kw):
    """手工构造 docmodel.Para，用于分类器的纯逻辑测试。"""
    from docmodel import Para
    defaults = dict(element=None, style_id=None, style_name=None,
                    outline_level=None, in_table=False, in_sdt=False,
                    section=0, has_field=False, pformat={})
    defaults.update(kw)
    return Para(index=index, text=text, **defaults)


def make_doc(paras):
    """把 Para 列表包成 classify() 能吃的最小 Doc。"""
    from docmodel import Doc
    return Doc(path="<memory>", paragraphs=paras, body_paragraphs=paras,
               sections=[], fields=[], stats={}, styles={})


def add_fake_image(p, anchor=False):
    """往段落里塞一个图形标记（假的 blip，不需要真图片数据）。

    检测只看元素名（drawing/blip/inline/anchor），构造一个最小结构就够，
    省得在测试里准备真图片文件。
    """
    WP = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
    r = p.add_run()
    drawing = r._element.makeelement(qn("w:drawing"), {})
    kind = drawing.makeelement(
        qn("wp:anchor" if anchor else "wp:inline"), {})
    kind.append(kind.makeelement(qn("a:blip"), {}))
    drawing.append(kind)
    r._element.append(drawing)
    return r


def para_has_image(p):
    return any(n.tag.split("}")[-1] in ("drawing", "blip", "pict", "object")
               for n in p.element.iter())
