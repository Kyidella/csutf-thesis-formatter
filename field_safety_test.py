"""域安全实验：验证不同的"改格式"方式会不会破坏 Zotero 引用域。

样本含 105 处 ADDIN ZOTERO_ITEM CSL_CITATION 域 + 1 个书目域。
本实验只在**确实含引用域的段落**上操作，并设置"危险做法"对照组。

用法:
    python field_safety_test.py <含引用域的.docx>
"""

import argparse
import copy
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Pt, Cm
from lxml import etree

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

TINY_PNG = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000a49444154789c6300010000050001" "0d0a2db40000000049454e44ae426082")
TINY_JPEG = bytes.fromhex(
    "ffd8ffe000104a46494600010101006000600000ffdb004300ffffffffffffffffff"
    "ffffffffffffffffffffffffffffffffffffffffffffffffffffffffffffffffffff"
    "ffffffffffffffffffffffffffffffffffc2000b080001000101011100ffc4001400"
    "01000000000000000000000000000000ffda0008010100003f00aaffd9")

TINY = {".png": TINY_PNG, ".jpeg": TINY_JPEG, ".jpg": TINY_JPEG, ".tiff": TINY_PNG}


def shrink(src, dst):
    """把样本里的图片全换成 1x1 占位图，便于快速反复实验（图片不影响域结构）。"""
    zin = zipfile.ZipFile(src)
    with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            n = item.filename.lower()
            if n.startswith("word/media/"):
                for ext, tiny in TINY.items():
                    if n.endswith(ext):
                        data = tiny
                        break
            zout.writestr(item, data)
    return Path(dst)


def field_report(path):
    """统计引用域数量与 fldChar 嵌套结构的完整性。"""
    z = zipfile.ZipFile(path)
    root = etree.fromstring(z.read("word/document.xml"))
    citations = bibl = orphan = unbalanced = depth = 0
    for r in root.iter(f"{W}r"):
        for child in r:
            tag = etree.QName(child).localname
            if tag == "fldChar":
                t = child.get(f"{W}fldCharType")
                if t == "begin":
                    depth += 1
                elif t == "end":
                    depth -= 1
                    if depth < 0:
                        unbalanced += 1
                        depth = 0
            elif tag == "instrText":
                txt = child.text or ""
                if "ZOTERO_ITEM" in txt:
                    if depth > 0:
                        citations += 1
                    else:
                        orphan += 1
                elif "ZOTERO_BIBL" in txt:
                    bibl += 1
    return {"citations": citations, "bibl": bibl, "orphan": orphan,
            "unbalanced": unbalanced, "unclosed": max(0, depth)}


def cite_paragraphs(docx_path):
    """返回含引用域的段落下标（按 python-docx 的段落索引）。"""
    z = zipfile.ZipFile(docx_path)
    root = etree.fromstring(z.read("word/document.xml"))
    body = root.find(f"{W}body")
    idxs = []
    for i, p in enumerate(body.findall(f"{W}p")):
        for it in p.iter(f"{W}instrText"):
            if "ZOTERO_ITEM" in (it.text or ""):
                idxs.append(i)
                break
    return idxs


def fmt_fingerprint(path, pidx=68):
    """目标段落的格式指纹：run 数、rPr 覆盖、斜体、字体直接覆盖、域元素。

    用来判断"改格式"有没有把段落原有的格式弄丢。
    """
    root = etree.fromstring(zipfile.ZipFile(path).read("word/document.xml"))
    p = root.find(f"{W}body").findall(f"{W}p")[pidx]
    runs = p.findall(f"{W}r")
    return {
        "runs": len(runs),
        "rPr": sum(1 for r in runs if r.find(f"{W}rPr") is not None),
        "斜体i": sum(1 for r in runs
                    if r.find(f"{W}rPr/{W}i") is not None),
        "直接字体rFonts": sum(1 for r in runs
                          if r.find(f"{W}rPr/{W}rFonts") is not None),
        "字号sz": sum(1 for r in runs
                     if r.find(f"{W}rPr/{W}sz") is not None),
        "instrText": sum(1 for r in runs if r.find(f"{W}instrText") is not None),
        "fldChar": sum(1 for r in runs if r.find(f"{W}fldChar") is not None),
    }


def xml_bytes(path):
    return zipfile.ZipFile(path).read("word/document.xml")


def run_case(name, base, outdir, mutate):
    """在 base 上执行 mutate(doc)，保存，返回 (名字, 域报告, XML 是否变化, 路径)。"""
    doc = Document(base)
    mutate(doc)
    out = Path(outdir) / f"{name}.docx"
    doc.save(out)
    changed = xml_bytes(out) != xml_bytes(base)
    return name, field_report(out), changed, out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("input")
    ap.add_argument("--outdir", default="exp")
    args = ap.parse_args()

    out = Path(args.outdir)
    out.mkdir(exist_ok=True)
    print(f"[1/4] 瘦身样本（{Path(args.input).stat().st_size/1024/1024:.0f}MB "
          f"→ 占位图）")
    base = shrink(args.input, out / "base_small.docx")
    baseline = field_report(base)
    print(f"      基线域状态: {baseline}")

    idxs = cite_paragraphs(base)
    print(f"      含引用域的段落: {len(idxs)} 个，前几个下标 {idxs[:6]}\n")
    if not idxs:
        print("样本里没找到引用域，实验无法进行"); return
    ti = idxs[0]
    base_fmt = fmt_fingerprint(base, ti)

    cases = []

    def case_a(doc):
        """安全做法：只改字号，不动字体（字体由样式/主题提供）。"""
        for r in doc.paragraphs[ti].runs:
            r.font.size = Pt(12)

    def case_a2(doc):
        """错误示范：把字体也一起写死成宋体。
        python-docx 的 run.font.name 写的是 w:ascii / w:hAnsi（西文字体），
        对中文文档来说是陷阱——会把英文的新罗马覆盖成宋体。"""
        for r in doc.paragraphs[ti].runs:
            r.font.name = "宋体"

    def case_b(doc):
        """中等风险：清掉所有 run 的直接格式。"""
        for r in doc.paragraphs[ti].runs:
            rpr = r._element.rPr
            if rpr is not None:
                r._element.remove(rpr)

    def case_c(doc):
        """文档级：改页边距，完全不碰正文。"""
        for sec in doc.sections:
            sec.top_margin = Cm(2.5)
            sec.left_margin = Cm(3.0)

    def case_d(doc):
        """危险做法：把段落里所有 run 合并成一个。"""
        p = doc.paragraphs[ti]
        runs = p.runs
        if not runs:
            return
        first = runs[0]._element
        for r in runs[1:]:
            for t in r._element.findall(f"{W}t"):
                first.append(copy.deepcopy(t))
            r._element.getparent().remove(r._element)

    def case_e(doc):
        """最危险：整段重写文本。"""
        p = doc.paragraphs[ti]
        p.text = p.text

    print("[2/4] 在同一个含引用域的段落上执行操作\n")
    for name, fn in (("A_只改字号", case_a),
                     ("A2_一起改字体", case_a2),
                     ("B_清run直接格式", case_b),
                     ("C_改页边距", case_c),
                     ("D_合并all_runs", case_d),
                     ("E_重写段落文本", case_e)):
        nm, rep, changed, path = run_case(name, base, out, fn)
        cases.append((nm, rep, changed, path))

    print("[3/4] 引用域完整性")
    print(f"      {'基线（未改动）':<20}: {baseline}")
    for nm, rep, changed, _ in cases:
        if not changed:
            verdict = "— 未产生实际改动"
        elif rep == baseline:
            verdict = "✓ 域完好"
        else:
            verdict = "✗ 域被破坏"
        print(f"      {nm:<20}: {rep}  {verdict}")

    print(f"\n[4/4] 目标段落格式指纹（基线: {base_fmt}）")
    for nm, _, _, path in cases:
        f = fmt_fingerprint(path, ti)
        lost = [k for k in ("斜体i", "直接字体rFonts", "rPr")
                if k in base_fmt and f[k] < base_fmt[k]]
        added = [k for k in ("直接字体rFonts", "字号sz")
                 if f[k] > base_fmt[k]]
        note = ""
        if lost:
            note += f"  丢失: {lost}"
        if added:
            note += f"  新增覆盖: {added}"
        print(f"      {nm:<20}: {f}{note}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
